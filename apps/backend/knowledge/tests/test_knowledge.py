"""Knowledge-base API, ingestion, chunking and retrieval tests.

Uses a deterministic bag-of-words embedder (no model download) and disables
the cross-encoder, so ranking is driven by the vector + BM25 signals alone.
BackgroundTasks run inside the ASGI call, so ingestion has finished by the
time the upload request returns.
"""
from __future__ import annotations

import hashlib
import io
import math
import re

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from src.api.deps import CurrentUser, get_current_user
from src.core.roles import Roles
from src.db.session import get_db

from knowledge.models import KnowledgeChunk, KnowledgeDocument
from knowledge.services import chunking, ingest, retrieval
from knowledge.services.document_parser import parse_document
from knowledge.services.document_safety import DocumentUploadError, validate_uploaded_document

BASE = "/api/v1/knowledge-bases"
_WORD = re.compile(r"[a-z0-9]+")


class FakeEmbedder:
    model_name = "fake-bow"
    max_input_tokens = 512
    dim = 64

    def count_tokens(self, text: str) -> int:
        return len(text.split())

    def _vec(self, text: str) -> list[float]:
        v = [0.0] * self.dim
        for w in _WORD.findall(text.lower()):
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % self.dim] += 1.0
        n = math.sqrt(sum(x * x for x in v)) or 1.0
        return [x / n for x in v]

    async def embed_texts(self, texts: list[str]) -> list[list[float]]:
        return [self._vec(t) for t in texts]

    async def embed_query(self, text: str) -> list[float]:
        return self._vec(text)


@pytest.fixture(autouse=True)
def _fake_models(monkeypatch, db: AsyncSession):
    fake = FakeEmbedder()
    for mod in (chunking, ingest, retrieval):
        monkeypatch.setattr(mod, "get_embedder", lambda: fake)
    monkeypatch.setattr(retrieval, "_rerank_scores", lambda query, texts: None)
    maker = async_sessionmaker(bind=db.bind, class_=AsyncSession, expire_on_commit=False)
    monkeypatch.setattr(ingest, "session_factory", lambda: maker())
    return fake


@pytest_asyncio.fixture
async def make_client(db: AsyncSession):
    """Build a client authenticated as an arbitrary ``CurrentUser``."""
    from apps.backend.main import create_app

    maker = async_sessionmaker(bind=db.bind, class_=AsyncSession, expire_on_commit=False)

    async def override_db():
        async with maker() as session:
            yield session

    clients: list[AsyncClient] = []

    async def build(user: CurrentUser) -> AsyncClient:
        app = create_app()
        app.dependency_overrides[get_current_user] = lambda: user
        app.dependency_overrides[get_db] = override_db
        client = AsyncClient(transport=ASGITransport(app=app), base_url="http://testserver")
        clients.append(client)
        return client

    yield build
    for c in clients:
        await c.aclose()


@pytest_asyncio.fixture
async def other_tenant(db: AsyncSession):
    from src.models import Tenant

    t = Tenant(id="tenant_test_02", name="Other", slug="other-tenant", status="active")
    db.add(t)
    await db.commit()
    return t


def _user(uid: str, tenant_id: str, role: str = Roles.TENANT_USER) -> CurrentUser:
    return CurrentUser(id=uid, email=f"{uid}@x.io", full_name=uid, role=role, tenant_id=tenant_id)


HANDBOOK = (
    "Refund policy. Customers can request a refund within 30 days of purchase. "
    "Refunds are paid back to the original card within five business days.\n\n"
    "Shipping. Orders ship from the Pune warehouse. Express shipping takes two days "
    "and standard shipping takes a week.\n\n"
    "Support hours. The support desk is open Monday to Friday, 9am to 6pm IST. "
    "Ticket IDs look like SUP-4821."
)


# ── API ──────────────────────────────────────────────────────────────────────


async def test_create_list_get(tenant_client, tenant_user):
    r = await tenant_client.post(BASE, json={"name": "  Support   docs ", "description": "FAQ"})
    assert r.status_code == 201, r.text
    kb = r.json()
    assert kb["name"] == "Support docs"
    assert kb["owner_id"] == tenant_user.id
    assert kb["document_count"] == 0 and kb["can_manage"] is True

    listed = (await tenant_client.get(BASE)).json()
    assert [k["id"] for k in listed] == [kb["id"]]
    assert (await tenant_client.get(f"{BASE}/{kb['id']}")).json()["name"] == "Support docs"


async def test_vendor_admin_has_no_tenant(admin_client):
    r = await admin_client.get(BASE)
    assert r.status_code == 403


async def test_other_tenant_cannot_see_or_touch(tenant_client, make_client, other_tenant):
    kb = (await tenant_client.post(BASE, json={"name": "Private"})).json()
    outsider = await make_client(_user("out_1", other_tenant.id, Roles.TENANT_ADMIN))
    assert (await outsider.get(BASE)).json() == []
    assert (await outsider.get(f"{BASE}/{kb['id']}")).status_code == 404
    assert (await outsider.delete(f"{BASE}/{kb['id']}")).status_code == 404
    r = await outsider.post(f"{BASE}/{kb['id']}/search", json={"query": "refund"})
    assert r.status_code == 404


async def test_manage_permissions(tenant_client, make_client, tenant):
    kb = (await tenant_client.post(BASE, json={"name": "Team KB"})).json()
    colleague = await make_client(_user("tu_2", tenant.id))
    admin = await make_client(_user("ta_1", tenant.id, Roles.TENANT_ADMIN))

    # A colleague can read and add, but not rename/delete someone else's KB.
    assert (await colleague.get(f"{BASE}/{kb['id']}")).json()["can_manage"] is False
    assert (await colleague.patch(f"{BASE}/{kb['id']}", json={"name": "x"})).status_code == 403
    assert (await colleague.delete(f"{BASE}/{kb['id']}")).status_code == 403
    doc = (await colleague.post(f"{BASE}/{kb['id']}/text", json={"title": "Note", "content": "hello"})).json()

    # The KB's creator can delete the colleague's document; the tenant admin can delete the KB.
    assert (await tenant_client.delete(f"{BASE}/{kb['id']}/documents/{doc['id']}")).status_code == 204
    assert (await admin.patch(f"{BASE}/{kb['id']}", json={"name": "Renamed"})).json()["name"] == "Renamed"
    assert (await admin.delete(f"{BASE}/{kb['id']}")).status_code == 204
    assert (await tenant_client.get(BASE)).json() == []


async def test_text_document_indexes_and_searches(tenant_client, db):
    kb = (await tenant_client.post(BASE, json={"name": "Handbook"})).json()
    r = await tenant_client.post(f"{BASE}/{kb['id']}/text", json={"title": "Customer handbook", "content": HANDBOOK})
    assert r.status_code == 201, r.text
    assert r.json()["status"] == "processing"

    docs = (await tenant_client.get(f"{BASE}/{kb['id']}/documents")).json()
    assert len(docs) == 1
    assert docs[0]["status"] == "ready", docs[0]
    assert docs[0]["editable"] is True and docs[0]["chunk_count"] >= 1
    assert docs[0]["filename"] == "Customer_handbook.md"

    stored = (await db.execute(select(func.count(KnowledgeChunk.id)))).scalar_one()
    assert stored == docs[0]["chunk_count"]

    hits = (await tenant_client.post(f"{BASE}/{kb['id']}/search", json={"query": "SUP-4821 support hours", "top_k": 3})).json()
    assert hits["results"], hits
    top = hits["results"][0]
    assert "SUP-4821" in top["content"]
    assert top["filename"] == "Customer_handbook.md"
    assert top["retrieval_method"] in {"hybrid", "keyword", "semantic"}
    assert (await tenant_client.get(f"{BASE}/{kb['id']}")).json()["document_count"] == 1


async def test_markdown_upload_is_editable_and_reindexes(tenant_client, db):
    kb = (await tenant_client.post(BASE, json={"name": "Docs"})).json()
    files = {"file": ("guide.md", b"# Setup\n\nInstall the agent with pip install widget.\n", "text/markdown")}
    doc = (await tenant_client.post(f"{BASE}/{kb['id']}/documents", files=files)).json()
    assert doc["editable"] is True

    text = (await tenant_client.get(f"{BASE}/{kb['id']}/documents/{doc['id']}/text")).json()
    assert text["title"] == "Setup"
    assert "pip install widget" in text["content"]

    r = await tenant_client.put(
        f"{BASE}/{kb['id']}/documents/{doc['id']}/text",
        json={"title": "Setup v2", "content": "Install the agent with the zebra installer."},
    )
    assert r.status_code == 200, r.text
    assert r.json()["filename"] == "Setup_v2.md"

    contents = (await db.execute(select(KnowledgeChunk.content))).scalars().all()
    assert contents and all("widget" not in c for c in contents)
    assert any("zebra" in c for c in contents)


async def test_binary_document_is_not_editable(tenant_client):
    kb = (await tenant_client.post(BASE, json={"name": "Sheets"})).json()
    files = {"file": ("prices.csv", b"item,price\napple,10\nmango,25\n", "text/csv")}
    doc = (await tenant_client.post(f"{BASE}/{kb['id']}/documents", files=files)).json()
    assert doc["editable"] is False
    docs = (await tenant_client.get(f"{BASE}/{kb['id']}/documents")).json()
    assert docs[0]["status"] == "ready"
    r = await tenant_client.get(f"{BASE}/{kb['id']}/documents/{doc['id']}/text")
    assert r.status_code == 409


async def test_unsupported_upload_rejected(tenant_client):
    kb = (await tenant_client.post(BASE, json={"name": "X"})).json()
    r = await tenant_client.post(f"{BASE}/{kb['id']}/documents", files={"file": ("run.exe", b"MZ\x90\x00", "application/octet-stream")})
    assert r.status_code == 415
    r = await tenant_client.post(f"{BASE}/{kb['id']}/documents", files={"file": ("fake.pdf", b"not a pdf", "application/pdf")})
    assert r.status_code == 415
    assert (await tenant_client.get(f"{BASE}/{kb['id']}/documents")).json() == []


async def test_failed_ingestion_is_recorded(tenant_client, monkeypatch):
    async def boom(*_a, **_k):
        raise RuntimeError("embedding provider down")

    monkeypatch.setattr(ingest, "chunk_document", boom)
    kb = (await tenant_client.post(BASE, json={"name": "X"})).json()
    await tenant_client.post(f"{BASE}/{kb['id']}/text", json={"title": "t", "content": "some text"})
    doc = (await tenant_client.get(f"{BASE}/{kb['id']}/documents")).json()[0]
    assert doc["status"] == "error"
    assert "embedding provider down" in doc["error_message"]


async def test_delete_document_and_kb_remove_chunks(tenant_client, db):
    kb = (await tenant_client.post(BASE, json={"name": "X"})).json()
    a = (await tenant_client.post(f"{BASE}/{kb['id']}/text", json={"title": "A", "content": HANDBOOK})).json()
    await tenant_client.post(f"{BASE}/{kb['id']}/text", json={"title": "B", "content": "Another note about mangoes."})

    assert (await tenant_client.delete(f"{BASE}/{kb['id']}/documents/{a['id']}")).status_code == 204
    left = (await db.execute(select(KnowledgeChunk.document_id).distinct())).scalars().all()
    assert a["id"] not in left and len(left) == 1

    assert (await tenant_client.delete(f"{BASE}/{kb['id']}")).status_code == 204
    assert (await db.execute(select(func.count(KnowledgeChunk.id)))).scalar_one() == 0
    assert (await db.execute(select(func.count(KnowledgeDocument.id)))).scalar_one() == 0


async def test_fail_interrupted_documents(tenant_client, db):
    kb = (await tenant_client.post(BASE, json={"name": "X"})).json()
    db.add(KnowledgeDocument(knowledge_base_id=kb["id"], tenant_id="tenant_test_01", owner_id="tu_1",
                             filename="stuck.pdf", status="processing"))
    await db.commit()
    assert await ingest.fail_interrupted_documents() == 1
    doc = (await tenant_client.get(f"{BASE}/{kb['id']}/documents")).json()[0]
    assert doc["status"] == "error" and "restart" in doc["error_message"]


# ── Parsing / chunking ───────────────────────────────────────────────────────


def test_parse_markdown_keeps_heading_path():
    parsed = parse_document("a.md", b"# Guide\n\n## Install\n\nRun the installer.\n")
    prose = [e for e in parsed.elements if e.kind == "prose"]
    assert prose[0].metadata["heading_path"] == "Guide > Install"


def test_parse_docx_and_xlsx():
    import docx
    import openpyxl

    d = docx.Document()
    d.add_heading("Leave policy", level=1)
    d.add_paragraph("Employees get 24 days of paid leave.")
    buf = io.BytesIO()
    d.save(buf)
    parsed = parse_document("p.docx", buf.getvalue())
    assert any("24 days" in e.text and e.metadata.get("heading_path") == "Leave policy" for e in parsed.elements)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["item", "price"])
    ws.append(["apple", 10])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_document("p.xlsx", buf.getvalue())
    table = parsed.elements[0]
    assert table.kind == "table" and table.metadata["sheet"] == "Prices" and "apple" in table.text


def test_safety_rejects_mismatched_and_empty_files():
    with pytest.raises(DocumentUploadError):
        validate_uploaded_document("x.docx", b"plain text")
    with pytest.raises(DocumentUploadError):
        validate_uploaded_document("x.txt", b"")
    with pytest.raises(DocumentUploadError):
        validate_uploaded_document("x.txt", b"bin\x00ary")


async def test_chunking_adds_context_and_respects_containers():
    parsed = parse_document("a.md", b"# Billing\n\nInvoices go out monthly.\n\n# Travel\n\nBook flights early.\n")
    chunks = await chunking.chunk_document(parsed, FakeEmbedder())
    assert len(chunks) == 2  # different headings never share a chunk
    assert chunks[0].content.startswith("Heading: Billing")
    assert chunks[1].metadata["heading_path"] == "Travel"
    assert chunks[0].metadata["embedding_model"] == "fake-bow"


def test_bm25_keeps_keyword_signal_in_a_tiny_knowledge_base():
    # Two chunks, the term in one of them: Okapi idf would be exactly 0 here.
    texts = ["Hotel limit in metro cities is 8000 INR.", "Employees get 24 days of leave."]
    assert retrieval._bm25_order("hotel limit", texts) == [0]
    assert retrieval._bm25_order("nothing matches", texts) == []
